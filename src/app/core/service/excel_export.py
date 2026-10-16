from __future__ import annotations

from abc import ABC, abstractmethod
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

from app.core.dto.reports.sales import SalesPeriodReport, MovementReport


class BaseExcelExporter(ABC):
    """
    Базовый экспортер в Excel.

    Задает шаблон: export() создает книгу, вызывает _build() у подкласса
    и возвращает байты .xlsx. Общие приемы оформления — здесь.
    """

    # --- Стили (общие для всех экспортеров) --------------------------
    TITLE_FONT = Font(bold=True, size=14)
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    DATE_FORMAT = "DD.MM.YYYY HH:MM"
    _BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Денежный формат: целые рубли, разделитель тысяч, знак ₽ в конце
    MONEY_FORMAT = '#,##0 "₽"'

    # Стили для окраски прибыли (стандартные цвета Эксель "хорошо/плохо")
    _POSITIVE_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    _POSITIVE_FONT = Font(color="000000")
    _NEGATIVE_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    _NEGATIVE_FONT = Font(color="000000")
    
    # --- Переопределяются в подклассах -------------------------------
    sheet_title: str = "Лист"
    report_title: str = "Отчет"

    def export(self, report) -> bytes:
        """
        Точка входа: строит книгу по отчету и возвращает байты .xlsx.

        При каждом вызове создается новая книга, поэтому экземпляр
        экспортера можно переиспользовать.
        """
        self._wb = Workbook()
        self._ws = self._wb.active
        self._ws.title = self.sheet_title

        self._build(report)

        return self._to_bytes()

    @abstractmethod
    def _build(self, report) -> None:
        """Наполняет лист данными конкретного отчета."""

    # --- Вспомогательные методы --------------------------------------

    def _write_title(self, period_start: datetime, period_end: datetime, start_col: int = 1) -> None:
        """Пишет название отчета и период, начиная со столбца start_col."""
        self._ws.cell(row=1, column=start_col, value=self.report_title).font = self.TITLE_FONT
        self._ws.cell(row=2, column=start_col, value=f"Период: {period_start:%d.%m.%Y} — {period_end:%d.%m.%Y}")

    def _style_header_row(self, row: int, col_count: int, start_col: int = 1) -> None:
        """Оформляет строку как заголовок таблицы."""
        for col in range(start_col, start_col + col_count):
            cell = self._ws.cell(row=row, column=col)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self._BORDER

    def _write_summary_block(self, start_row: int, rows: list[tuple[str, int]], start_col: int = 1) -> int:
        """Пишет блок сводных показателей, начиная со столбца start_col."""
        self._ws.cell(row=start_row, column=start_col, value="Показатель")
        self._ws.cell(row=start_row, column=start_col + 1, value="Значение")
        self._style_header_row(start_row, 2, start_col)

        row = start_row
        for label, value in rows:
            row += 1
            self._ws.cell(row=row, column=start_col, value=label)
            self._ws.cell(row=row, column=start_col + 1, value=value)

        self._apply_table_borders(start_row + 1, row, 2, start_col)
        return row

    def _apply_table_borders(self, min_row: int, max_row: int, col_count: int, start_col: int = 1) -> None:
        """Обводит рамками блок данных."""
        for r in range(min_row, max_row + 1):
            for c in range(start_col, start_col + col_count):
                self._ws.cell(row=r, column=c).border = self._BORDER

    def _autofit_columns(self) -> None:
        """Подбирает ширину колонок по содержимому."""
        for col_cells in self._ws.columns:
            max_len = 0
            for cell in col_cells:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            letter = get_column_letter(col_cells[0].column)
            self._ws.column_dimensions[letter].width = min(max_len + 4, 60)

    def _set_auto_filter(self, header_row: int, last_row: int, col_count: int) -> None:
        """
        Включает автофильтр Excel для таблицы.

        Диапазон задаётся от строки заголовка до последней строки данных,
        поэтому стрелки фильтров появляются в шапке сразу при открытии файла.
        """
        if last_row < header_row:
            return
        last_col = get_column_letter(col_count)
        self._ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
    
    def _colorize_profit_column(self, col_index: int, data_start: int, last_row: int) -> None:
        """
        Окрашивает значения прибыли условным форматированием:
        положительные — в зелёный, отрицательные — в красный.
        """
        if last_row < data_start:
            return
        col_letter = get_column_letter(col_index)
        cell_range = f"{col_letter}{data_start}:{col_letter}{last_row}"

        self._ws.conditional_formatting.add(
            cell_range,
            CellIsRule(operator="greaterThan", formula=["0"],
                       fill=self._POSITIVE_FILL, font=self._POSITIVE_FONT),
        )
        self._ws.conditional_formatting.add(
            cell_range,
            CellIsRule(operator="lessThan", formula=["0"],
                       fill=self._NEGATIVE_FILL, font=self._NEGATIVE_FONT),
        )
    
    def _to_bytes(self) -> bytes:
        """Сохраняет книгу в память и возвращает байты."""
        buffer = BytesIO()
        self._wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    
    @staticmethod
    def _naive_datetime(dt: datetime | None) -> datetime | None:
        """
        Убирает информацию о таймзоне из datetime.

        openpyxl не поддерживает даты с таймзонами, поэтому перед
        записью в ячейку любой tz-aware datetime нужно сделать наивным.
        """
        if dt is None:
            return None
        return dt.replace(tzinfo=None)


class SalesReportExcelExporter(BaseExcelExporter):
    """Экспорт отчета по продажам."""

    sheet_title = "Продажи"
    report_title = "Отчет по продажам"

    def _build(self, report: SalesPeriodReport) -> None:
        self._write_title(report.period_start, report.period_end)

        summary_start_row = 4

        summary = [
            ("Продано товаров", report.total_sold_count),
            ("Выручка", report.total_revenue),
            ("Себестоимость", report.total_purchase_cost),
            ("Прибыль", report.total_profit),
            ("Средняя прибыль на товар", report.avg_profit_per_item),
        ]

        row = self._write_summary_block(summary_start_row, summary)

        # Финансовое форматирование для сводных показателей.
        # Строки:
        # summary_start_row     — заголовок "Показатель / Значение"
        # summary_start_row + 1 — Продано товаров, не деньги
        # summary_start_row + 2 — Выручка
        # summary_start_row + 3 — Себестоимость
        # summary_start_row + 4 — Прибыль
        # summary_start_row + 5 — Средняя прибыль на товар
        for offset in (2, 3, 4, 5):
            self._ws.cell(
                row=summary_start_row + offset,
                column=2,
            ).number_format = self.MONEY_FORMAT

        # Окрашиваем прибыль в сводке:
        # положительная — зелёным, отрицательная — красным.
        profit_summary_row = summary_start_row + 4
        self._colorize_profit_column(
            col_index=2,
            data_start=profit_summary_row,
            last_row=profit_summary_row,
        )

        # Детализация по играм
        row += 2

        headers = [
            "Игра",
            "Продано, шт",
            "Выручка",
            "Прибыль",
            "Средняя цена за шт",
        ]

        for col, header in enumerate(headers, start=1):
            self._ws.cell(row=row, column=col, value=header)

        self._style_header_row(row, len(headers))
        header_row = row

        # Колонки с деньгами:
        # 3 — Выручка
        # 4 — Прибыль
        # 5 — Средняя цена за шт
        MONEY_COLS = (3, 4, 5)

        # Колонка с прибылью для условного окрашивания
        PROFIT_COL = 4

        data_start = row + 1

        for item in report.sales_by_game:
            row += 1

            self._ws.cell(row=row, column=1, value=item.game_name)
            self._ws.cell(row=row, column=2, value=item.sold_count)
            self._ws.cell(row=row, column=3, value=item.total_revenue)
            self._ws.cell(row=row, column=4, value=item.total_profit)
            self._ws.cell(row=row, column=5, value=item.avg_price_per_unit)

        # Применяем денежный формат ко всем финансовым колонкам таблицы
        for r in range(data_start, row + 1):
            for money_col in MONEY_COLS:
                self._ws.cell(row=r, column=money_col).number_format = self.MONEY_FORMAT

        self._set_auto_filter(header_row, row, len(headers))

        # Окрашиваем прибыль в детализации по играм
        self._colorize_profit_column(
            col_index=PROFIT_COL,
            data_start=data_start,
            last_row=row,
        )

        self._apply_table_borders(data_start, row, len(headers))
        self._autofit_columns()


class MovementReportExcelExporter(BaseExcelExporter):
    """Экспорт отчета по движению товаров."""

    sheet_title = "Движение"
    report_title = "Отчет по движению товаров"

    def _build(self, report: MovementReport) -> None:
        # Заголовок, период и сводку сдвигаем на 1 столбец вправо (начиная с B),
        # чтобы колонка A с ID не растягивалась под длинный заголовок
        self._write_title(report.period_start, report.period_end, start_col=2)

        summary = [
            ("Всего записей", report.total_purchases),
            ("Продано", report.sold_count),
            ("В наличии", report.available_count),
            ("В пути к покупателю", report.in_transit_to_client_count),
            ("В пути ко мне", report.in_transit_to_me_count),
        ]
        row = self._write_summary_block(start_row=4, rows=summary, start_col=2)

        # Таблица позиций — по-прежнему с колонки A
        row += 2
        headers = [
            "ID", "Игра", "Транзакция", "Дата покупки", "Дата продажи",
            "Цена покупки", "Цена продажи", "Продано за", "Прибыль",
            "Статус", "Обмен", "Резерв", "Комментарий",
        ]
        for col, header in enumerate(headers, start=1):
            self._ws.cell(row=row, column=col, value=header)
        self._style_header_row(row, len(headers))
        header_row = row

        PROFIT_COL = 9
        MONEY_COLS = (6, 7, 8, PROFIT_COL)

        data_start = row + 1
        for item in report.items:
            row += 1
            self._ws.cell(row=row, column=1, value=item.id)
            self._ws.cell(row=row, column=2, value=item.game_name)
            self._ws.cell(row=row, column=3, value=item.utility_title)
            self._ws.cell(row=row, column=4, value=self._naive_datetime(item.purchase_at))
            self._ws.cell(row=row, column=5, value=self._naive_datetime(item.sold_at))
            self._ws.cell(row=row, column=6, value=item.price_purchase)
            self._ws.cell(row=row, column=7, value=item.price_selling)
            self._ws.cell(row=row, column=8, value=item.price_sold)
            self._ws.cell(row=row, column=9, value=item.profit)
            self._ws.cell(row=row, column=10, value=item.status_display)
            self._ws.cell(row=row, column=11, value=item.swap)
            self._ws.cell(row=row, column=12, value=item.reserve)
            self._ws.cell(row=row, column=13, value=item.comment)
            

        # Форматы ячеек: даты и деньги (со знаком рубля)
        for r in range(data_start, row + 1):
            self._ws.cell(row=r, column=4).number_format = self.DATE_FORMAT
            self._ws.cell(row=r, column=5).number_format = self.DATE_FORMAT
            for money_col in MONEY_COLS:
                self._ws.cell(row=r, column=money_col).number_format = self.MONEY_FORMAT

        self._set_auto_filter(header_row, row, len(headers))
        self._colorize_profit_column(PROFIT_COL, data_start, row)
        self._apply_table_borders(data_start, row, len(headers))
        self._autofit_columns()